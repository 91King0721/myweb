"""Import the six-block computing-center timetable (requires openpyxl).

Usage: python scripts/import-computing-center.py path/to/timetable.xlsx
Only room occupancy is published; course and teacher names stay in the source.
"""

import argparse
import hashlib
import json
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl

BUILDING = "计网中心"
BLOCKS = ["第一大节", "第二大节", "第三大节", "第四大节", "第五大节", "第六大节"]


def parse_weeks(value, limit):
    weeks = set()
    for part in value.split("、"):
        match = re.fullmatch(r"(\d+)(?:[–—-](\d+))?", part)
        if not match:
            raise ValueError(f"Unsupported weeks: {value}")
        first, last = int(match[1]), int(match[2] or match[1])
        if not 1 <= first <= last <= limit:
            raise ValueError(f"Weeks out of range: {value}")
        weeks.update(range(first, last + 1))
    return weeks


def import_timetable(source, target):
    text = target.read_text(encoding="utf-8")
    meta_match = re.search(r"^var DATA_META = (.*);$", text, re.M)
    data_match = re.search(r"^var DATA = (.*);$", text, re.M)
    meta, data = json.loads(meta_match[1]), json.loads(data_match[1])
    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook["2026秋课表"]
    if [sheet.cell(3, col).value for col in range(2, 9)] != ["星期" + d for d in "一二三四五六日"]:
        raise ValueError("Unexpected weekday headers")
    entries = []
    for block, label in enumerate(BLOCKS):
        row = block + 4
        if sheet.cell(row, 1).value != label:
            raise ValueError(f"Unexpected period label at row {row}")
        for day in range(7):
            cell = sheet.cell(row, day + 2)
            for line in (cell.value or "").splitlines():
                if not line.strip():
                    continue
                match = re.fullmatch(r".+?\s{2,}([\d、–—-]+)周\s{2,}(\d{3})(?:\s{2,}.*)?", line.strip())
                if not match:
                    raise ValueError(f"Unparsed entry at {cell.coordinate}: {line}")
                entries.append((parse_weeks(match[1], meta["weekCount"]), day, block, match[2]))
    workbook.close()
    room_names = sorted({entry[3] for entry in entries})
    if not room_names:
        raise ValueError("No rooms found")
    for week in range(1, meta["weekCount"] + 1):
        days = {}
        for day, day_name in enumerate(meta["dayNames"]):
            rooms = []
            for name in room_names:
                periods = [0] * 12
                for weeks, entry_day, block, room in entries:
                    if week in weeks and entry_day == day and room == name:
                        periods[block * 2:block * 2 + 2] = [1, 1]
                rooms.append({"name": f"{BUILDING}-{name}", "seats": "—", "periods": periods})
            current = date.fromisoformat(meta["week1Start"]) + timedelta(days=(week - 1) * 7 + day)
            days[day_name] = {
                "date": f"{current.month}.{current.day}", "isoDate": current.isoformat(),
                "free": sum(not any(room["periods"]) for room in rooms),
                "total": len(rooms), "rooms": rooms,
            }
        data[str(week)][BUILDING] = days
    if BUILDING not in meta["buildings"]:
        meta["buildings"].append(BUILDING)
    meta.setdefault("supplementalSources", {})[BUILDING] = {
        "file": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sheet": sheet.title, "range": "B4:H9", "courseEntries": len(entries),
        "periodMapping": [[1, 2], [3, 4], [5, 6], [7, 8], [9, 10], [11, 12]],
        "availability": "按课表周次推算，未排课节次记为空闲；不代表实际开放或无临时借用。",
        "seats": "课表未提供，显示为—，不按容量排除。",
    }
    compact = lambda value: json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    text = text[:data_match.start(1)] + compact(data) + text[data_match.end(1):]
    text = text[:meta_match.start(1)] + compact(meta) + text[meta_match.end(1):]
    target.write_text(text, encoding="utf-8", newline="\n")
    print(f"Imported {len(entries)} entries, {len(room_names)} rooms, {meta['weekCount']} weeks: {', '.join(room_names)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    args = parser.parse_args()
    import_timetable(args.source, Path(__file__).resolve().parents[1] / "assets/js/data.js")
